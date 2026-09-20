from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document

PACKAGE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PACKAGE_ROOT / "src"))

from paper4_pipeline.agents.protocols import AgentCallMetadata  # noqa: E402
from paper4_pipeline.agents.prompts import load_prompt  # noqa: E402
from paper4_pipeline.domain.models import (  # noqa: E402
    LearningObjective, ModelConfig, ProcedureStep, TeachingArtifact, TokenUsage,
)
from paper4_pipeline.providers.openai_compatible import (  # noqa: E402
    OpenAICompatibleProvider, ProviderCallResult, ProviderInvocationError,
)
from paper4_pipeline.web_api.docx_ingestion import extract_docx  # noqa: E402
from paper4_pipeline.web_api.schemas import (  # noqa: E402
    EngineLessonInput, EngineMode, RawBlock, RawLessonDocument,
)
from paper4_pipeline.web_api.segmented_normalizer import (  # noqa: E402
    ActivityOutput, OverviewOutput, SegmentedDocxNormalizer, _chunks,
    _merge_overviews,
)
from paper4_pipeline.web_api.settings import EngineSettings  # noqa: E402


class StubProvider:
    def __init__(self) -> None:
        self.settings = SimpleNamespace(provider="stub", model_name="stub")
        self.calls: list[str] = []

    def invoke_structured(self, *, prompt, input_payload, output_schema, stage):
        self.calls.append(stage)
        if stage == "docx_overview":
            value = OverviewOutput(
                content_analysis="原稿分析",
                learning_objectives=[LearningObjective(
                    objective_id="goal-1", description="识别关键线索",
                )],
            )
        else:
            index = input_payload["batch_index"]
            value = ActivityOutput(
                teaching_artifacts=[TeachingArtifact(
                    artifact_id="local-artifact", artifact_type="文本",
                    title=f"材料{index}", content=f"原文片段{index}",
                )],
                procedure_steps=[ProcedureStep(
                    step_id="local-step", stage=f"活动{index}",
                    duration_minutes=10, objective_ids=["goal-1"],
                    artifact_ids=["local-artifact"],
                    student_actions=[f"完成任务{index}"],
                )],
            )
        return ProviderCallResult(
            value=value, usage=TokenUsage(input_tokens=100, output_tokens=50),
            estimated_cost=0.01,
            metadata=AgentCallMetadata(
                provider="stub", model_name="stub", prompt_id=prompt.prompt_id,
                prompt_version=prompt.version, prompt_sha256=prompt.sha256,
                attempts=1,
            ),
        )


class SegmentedImportTests(unittest.TestCase):
    def test_multi_part_overview_keeps_both_objectives_with_unique_ids(self) -> None:
        merged = _merge_overviews([
            OverviewOutput(
                content_analysis="片段一",
                learning_objectives=[LearningObjective(
                    objective_id="goal", description="识别线索",
                )],
            ),
            OverviewOutput(
                content_analysis="片段二",
                learning_objectives=[LearningObjective(
                    objective_id="goal", description="分析情感",
                )],
            ),
        ])
        self.assertIn("片段一", merged.content_analysis)
        self.assertIn("片段二", merged.content_analysis)
        self.assertEqual(2, len({item.objective_id for item in merged.learning_objectives}))

    def test_truncated_import_stops_after_one_call_and_records_exception_usage(self) -> None:
        provider = OpenAICompatibleProvider(ModelConfig(max_retries=2))
        client = Mock()
        client.bind.return_value.invoke.side_effect = RuntimeError(
            "length limit was reached - CompletionUsage("
            "completion_tokens=8192, prompt_tokens=14000, total_tokens=22192)"
        )
        provider._client = client
        with self.assertRaises(ProviderInvocationError) as caught:
            provider.invoke_structured(
                prompt=load_prompt("docx_overview_prompt"),
                input_payload={"source_blocks": []},
                output_schema=OverviewOutput, stage="docx_overview",
            )
        self.assertEqual(1, caught.exception.attempts)
        self.assertEqual(14000, caught.exception.usage.input_tokens)
        self.assertEqual(8192, caught.exception.usage.output_tokens)
        self.assertEqual(1, client.bind.return_value.invoke.call_count)

    def test_docx_reader_preserves_paragraph_table_paragraph_order(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lesson.docx"
            document = Document()
            document.add_paragraph("开场")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "教师提问"
            table.cell(0, 1).text = "学生回答"
            document.add_paragraph("小结")
            document.save(path)
            settings = EngineSettings.from_environment()
            raw = extract_docx(path, original_filename=path.name, settings=settings)
            self.assertEqual(
                ["paragraph", "table_row", "paragraph"],
                [block.kind for block in raw.ordered_blocks],
            )
            self.assertEqual("开场", raw.ordered_blocks[0].text)
            self.assertIn("学生回答", raw.ordered_blocks[1].text)

    def test_external_word_is_imported_in_bounded_chunks_with_traceable_content(self) -> None:
        blocks = [
            RawBlock(locator="paragraph:1", kind="paragraph", text="活动A" * 500),
            RawBlock(locator="paragraph:2", kind="paragraph", text="活动B" * 500),
        ]
        raw = RawLessonDocument(
            source_sha256="a" * 64, original_filename="external.docx",
            ordered_blocks=blocks,
        )
        self.assertEqual(2, len(_chunks(blocks)))
        provider = StubProvider()
        task = EngineLessonInput(
            mode=EngineMode.OPTIMIZE, subject="语文", grade="八年级",
            topic="《背影》", duration_minutes=45,
            must_preserve_content=["活动A"],
        )
        normalized = SegmentedDocxNormalizer(provider).normalize(
            raw, task, task_id="task-1", plan_id="plan-1",
        )
        self.assertEqual(["docx_overview", "docx_activities", "docx_activities"], provider.calls)
        self.assertEqual(2, len(normalized.lesson_plan.procedure_steps))
        self.assertEqual(2, len(normalized.lesson_plan.teaching_artifacts))
        self.assertEqual(45, sum(x.duration_minutes for x in normalized.lesson_plan.procedure_steps))
        self.assertEqual("活动2", normalized.lesson_plan.procedure_steps[1].stage)
        self.assertEqual(3, normalized.model_metadata["attempts"])
        self.assertEqual(300, normalized.usage.input_tokens)
        self.assertEqual(150, normalized.usage.output_tokens)
        self.assertTrue(any("人工核对" in item for item in normalized.warnings))


if __name__ == "__main__":
    unittest.main()
