from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from concurrent.futures import Executor, Future
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from fastapi.testclient import TestClient


PACKAGE_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PACKAGE_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from paper4_pipeline.web_api.app import create_app  # noqa: E402
from paper4_pipeline.web_api.docx_ingestion import (  # noqa: E402
    DocxIngestionError,
    extract_docx,
    inspect_docx,
)
from paper4_pipeline.web_api.manager import EngineManager  # noqa: E402
from paper4_pipeline.web_api.generated_docx import restore_generated_docx  # noqa: E402
from paper4_pipeline.web_api.normalizer import DocxNormalizer  # noqa: E402
from paper4_pipeline.web_api.projection import EngineProjection  # noqa: E402
from paper4_pipeline.web_api.registry import RunRecord, RunRegistry  # noqa: E402
from paper4_pipeline.web_api.runner import (  # noqa: E402
    _classify_error, _fail, build_task, _public_message, execute_optimize,
)
from paper4_pipeline.web_api.schemas import (  # noqa: E402
    CreateRunRequest,
    EngineLessonInput,
    EngineMode,
    EngineRunStatus,
    NormalizationModelOutput,
    NormalizedLessonInput,
    RawLessonDocument,
    RunAccepted,
)
from paper4_pipeline.web_api.settings import EngineSettings  # noqa: E402
from paper4_pipeline.agents.protocols import AgentCallMetadata  # noqa: E402
from paper4_pipeline.domain.models import (  # noqa: E402
    PipelineResult,
    RunStatus,
    StopReason,
    TokenUsage,
)
from paper4_pipeline.providers.openai_compatible import (  # noqa: E402
    ProviderCallResult,
    ProviderInvocationError,
)
from paper4_pipeline.exporters.common import sha256_file  # noqa: E402
from support import make_config, make_document, make_task  # noqa: E402


JOB_ID = "01ARZ3NDEKTSV4RRFFQ69G5FAV"
REQUEST_HASH = "a" * 64


class RecordingExecutor(Executor):
    def __init__(self) -> None:
        self.calls: list[tuple[object, tuple[object, ...]]] = []

    def submit(self, fn, /, *args, **kwargs):  # type: ignore[no-untyped-def]
        self.calls.append((fn, args))
        future: Future[None] = Future()
        future.set_result(None)
        return future


def make_settings(root: Path, **updates: object) -> EngineSettings:
    values: dict[str, object] = {
        "state_root": root / "state",
        "artifacts_root": root / "artifacts",
        "config_path": PACKAGE_ROOT / "configs" / "deepseek_v4_flash.json",
        "internal_token": "test-token",
    }
    values.update(updates)
    return EngineSettings(**values)


def request_payload(*, mode: str = "generate") -> dict[str, object]:
    return {
        "contract_version": "1",
        "external_job_id": JOB_ID,
        "request_sha256": REQUEST_HASH,
        "task": {
            "mode": mode,
            "subject": "数学",
            "grade": "高一",
            "topic": "函数的单调性",
            "duration_minutes": 45,
        },
    }


class RunRegistryTests(unittest.TestCase):
    def test_create_is_persistent_and_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            settings = make_settings(Path(directory))
            registry = RunRegistry(settings)
            record = RunRecord(
                engine_run_id="run-one",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.GENERATE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
            )
            created, is_new = registry.create(record)
            same, second_is_new = RunRegistry(settings).create(record)
            self.assertTrue(is_new)
            self.assertFalse(second_is_new)
            self.assertEqual(created.engine_run_id, same.engine_run_id)

    def test_reusing_job_id_with_other_payload_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            registry = RunRegistry(make_settings(Path(directory)))
            base = RunRecord(
                engine_run_id="run-one",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.GENERATE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
            )
            registry.create(base)
            with self.assertRaisesRegex(ValueError, "IDEMPOTENCY"):
                registry.create(
                    base.model_copy(
                        update={"engine_run_id": "run-two", "request_sha256": "b" * 64}
                    )
                )

    def test_restart_marks_orphaned_active_run_failed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            registry = RunRegistry(make_settings(Path(directory)))
            record = RunRecord(
                engine_run_id="run-interrupted",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.GENERATE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
            )
            registry.create(record)
            self.assertEqual(1, registry.recover_interrupted_runs())
            recovered = registry.get(record.engine_run_id)
            self.assertEqual("failed", recovered.status.value)
            self.assertEqual("ENGINE_RESTART_INTERRUPTED", recovered.error_code)

    def test_restart_recovers_a_durable_terminal_result(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            settings = make_settings(Path(directory))
            registry = RunRegistry(settings)
            record = RunRecord(
                engine_run_id="run-finished-before-registry-update",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.GENERATE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
                status=EngineRunStatus.RUNNING,
            )
            registry.create(record)
            artifact_dir = settings.artifacts_root / record.engine_run_id
            artifact_dir.mkdir(parents=True)
            # An empty completed result is not a normal pipeline output, but it is
            # sufficient here to freeze the registry's status-recovery contract.
            result = PipelineResult(
                run_id=record.engine_run_id,
                task_id=record.task_id,
                experiment_id="recovery-test",
                method_id="durable-result-fixture",
                status=RunStatus.COMPLETED,
                stop_reason=StopReason.QUALITY_PASSED,
            )
            (artifact_dir / "run_result.json").write_text(
                result.model_dump_json(indent=2), encoding="utf-8"
            )

            self.assertEqual(1, registry.recover_interrupted_runs())

            recovered = registry.get(record.engine_run_id)
            self.assertEqual(EngineRunStatus.COMPLETED, recovered.status)
            self.assertEqual("quality_passed", recovered.stop_reason)

    def test_stage_changes_have_unique_monotonic_lifecycle_sequences(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            registry = RunRegistry(make_settings(Path(directory)))
            record = RunRecord(
                engine_run_id="run-events",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.OPTIMIZE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
            )
            registry.create(record)
            registry.update(record.engine_run_id, status=EngineRunStatus.PREPROCESSING, stage="docx_security_check")
            registry.update(record.engine_run_id, stage="docx_normalize")
            registry.update(record.engine_run_id, status=EngineRunStatus.RUNNING, stage="judge")
            sequences = [item["sequence"] for item in registry.lifecycle_events(record.engine_run_id)]
            self.assertEqual([1, 2, 3, 4], sequences)
            self.assertEqual(len(sequences), len(set(sequences)))


class DocxIngestionTests(unittest.TestCase):
    def test_extracts_paragraphs_and_tables_with_stable_locators(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "lesson.docx"
            document = Document()
            document.add_heading("一次函数", level=1)
            document.add_paragraph("教学目标：理解斜率的意义")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "环节"
            table.cell(0, 1).text = "学生活动"
            table.cell(1, 0).text = "探究"
            table.cell(1, 1).text = "绘制图像并说明变化"
            document.save(path)

            raw = extract_docx(
                path,
                original_filename="我的教案.docx",
                settings=make_settings(root),
            )
            self.assertEqual("paragraph:1", raw.paragraphs[0].locator)
            self.assertEqual("table:1/row:1/column:1", raw.tables[0].cells[0].locator)
            self.assertIn("绘制图像", raw.tables[0].cells[-1].text)

    def test_merged_table_cell_is_extracted_once_not_once_per_grid_column(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "merged-template.docx"
            document = Document()
            table = document.add_table(rows=2, cols=4)
            merged = table.cell(0, 0).merge(table.cell(0, 3))
            merged.text = "一、课程标准：理解并运用比较级"
            table.cell(1, 0).text = "环节"
            table.cell(1, 1).text = "教师活动"
            table.cell(1, 2).text = "学生活动"
            table.cell(1, 3).text = "评价"
            document.save(path)

            raw = extract_docx(
                path, original_filename=path.name, settings=make_settings(root),
            )
            matching = [
                cell for cell in raw.tables[0].cells
                if "课程标准" in cell.text
            ]
            self.assertEqual(1, len(matching))
            self.assertEqual(1, raw.ordered_blocks[0].text.count("课程标准"))

    def test_rejects_zip_slip_member(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "bad.docx"
            with zipfile.ZipFile(path, "w") as archive:
                for name in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                    archive.writestr(name, "<root />")
                archive.writestr("../escape.txt", "bad")
            with self.assertRaisesRegex(DocxIngestionError, "不安全路径"):
                inspect_docx(path, make_settings(root))

    def test_accepts_external_relationship_without_dereferencing_it(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "external.docx"
            relations = (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Target="https://example.test/template" '
                'TargetMode="External" Type="template" /></Relationships>'
            )
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("[Content_Types].xml", "<root />")
                archive.writestr("_rels/.rels", relations)
                archive.writestr("word/document.xml", "<root />")
            warnings = inspect_docx(path, make_settings(root))
            self.assertTrue(any("外部关系" in item for item in warnings))
            self.assertTrue(any("不会访问" in item for item in warnings))

    def test_external_images_and_templates_are_warnings_not_rejections(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "linked-assets.docx"
            relations = (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Target="file:///D:/missing/header.png" '
                'TargetMode="External" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" />'
                '<Relationship Id="rId2" Target="file:///D:/missing/template.dotx" '
                'TargetMode="External" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" />'
                '</Relationships>'
            )
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("[Content_Types].xml", "<root />")
                archive.writestr("_rels/.rels", relations)
                archive.writestr("word/document.xml", "<root />")
            warnings = inspect_docx(path, make_settings(root))
            self.assertTrue(any("外链图片" in item for item in warnings))
            self.assertTrue(any("外部模板" in item for item in warnings))

    def test_rejects_case_variant_macro_part(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "macro.docx"
            with zipfile.ZipFile(path, "w") as archive:
                for name in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                    archive.writestr(name, "<root />")
                archive.writestr("word/VBAPROJECT.BIN", b"macro")
            with self.assertRaisesRegex(DocxIngestionError, "宏"):
                inspect_docx(path, make_settings(root))

    def test_rejects_fake_extension_and_corrupt_zip(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fake = root / "fake.docx"
            fake.write_bytes(b"this is not a Word package")
            with self.assertRaisesRegex(DocxIngestionError, "OOXML"):
                inspect_docx(fake, make_settings(root))

            corrupt = root / "corrupt.docx"
            corrupt.write_bytes(b"PK\x03\x04broken")
            with self.assertRaisesRegex(DocxIngestionError, "损坏"):
                inspect_docx(corrupt, make_settings(root))

    def test_rejects_ole_encrypted_and_high_compression_packages(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for filename, member, expected in (
                ("ole.docx", "word/embeddings/object1.bin", "OLE"),
                ("encrypted.docx", "EncryptedPackage", "加密"),
            ):
                path = root / filename
                with zipfile.ZipFile(path, "w") as archive:
                    for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                        archive.writestr(required, "<root />")
                    archive.writestr(member, b"unsafe")
                with self.assertRaisesRegex(DocxIngestionError, expected):
                    inspect_docx(path, make_settings(root))

            bomb = root / "bomb.docx"
            with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                    archive.writestr(required, "<root />")
                archive.writestr("word/large.txt", b"0" * (2 * 1024 * 1024))
            with self.assertRaisesRegex(DocxIngestionError, "压缩比"):
                inspect_docx(bomb, make_settings(root, max_compression_ratio=10.0))

    def test_rejects_upload_and_entry_count_limits(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            oversized = root / "oversized.docx"
            oversized.write_bytes(b"PK\x03\x04" + b"0" * 64)
            with self.assertRaisesRegex(DocxIngestionError, "20 MiB"):
                inspect_docx(oversized, make_settings(root, max_upload_bytes=16))

            entries = root / "entries.docx"
            with zipfile.ZipFile(entries, "w") as archive:
                for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                    archive.writestr(required, "<root />")
                archive.writestr("word/extra.xml", "<root />")
            with self.assertRaisesRegex(DocxIngestionError, "过多"):
                inspect_docx(entries, make_settings(root, max_zip_entries=3))


class TaskMappingTests(unittest.TestCase):
    def test_generate_input_maps_to_real_lesson_task(self) -> None:
        task_input = EngineLessonInput(
            mode=EngineMode.GENERATE,
            subject="物理",
            grade="高二",
            topic="动量守恒",
            class_size=48,
            textbook_version="人教版",
        )
        task = build_task(task_input, task_id="lesson-test")
        self.assertEqual("generate", task.mode.value)
        self.assertEqual(48, task.class_constraints["class_size"])
        self.assertEqual("人教版", task.metadata["textbook_version"])
        self.assertIn("procedure_steps", task.required_sections)

    def test_empty_key_does_not_corrupt_public_error_message(self) -> None:
        with patch.dict("os.environ", {"DEEPSEEK_API_KEY": ""}):
            self.assertEqual("plain failure", _public_message("plain failure"))

    def test_docx_output_limit_has_actionable_public_error(self) -> None:
        error = RuntimeError(
            "docx_normalizer failed: LengthFinishReasonError: "
            "Could not parse response content as the length limit was reached"
        )
        self.assertEqual("DOCX_NORMALIZATION_TOO_LARGE", _classify_error(error))
        self.assertIn("结构化服务的输出被截断", _public_message(str(error), code=_classify_error(error)))


class DocxNormalizerAccountingTests(unittest.TestCase):
    def test_empty_optional_textbook_is_not_authoritative(self) -> None:
        task = make_task()
        document = make_document(task, high_quality=True)
        document = document.model_copy(update={
            "metadata": document.metadata.model_copy(update={"textbook_version": "人教版"})
        })
        task_input = EngineLessonInput(
            mode=EngineMode.OPTIMIZE,
            subject=task.subject,
            grade=task.grade,
            topic=task.topic,
            duration_minutes=task.duration_minutes,
        )
        normalizer = DocxNormalizer(make_config())
        normalizer.provider = Mock()
        normalizer.provider.invoke_structured.return_value = ProviderCallResult(
            value=NormalizationModelOutput(
                metadata_provenance={
                    "subject": "user", "grade": "user", "topic": "user",
                    "duration_minutes": "user", "textbook_version": "file",
                },
                lesson_plan=document,
            ),
            usage=TokenUsage(),
            estimated_cost=0,
            metadata=AgentCallMetadata(
                provider="test", model_name="test", prompt_id="docx_normalizer_prompt",
                prompt_version="1.1", prompt_sha256="a" * 64,
            ),
        )
        normalizer.normalize(
            RawLessonDocument(source_sha256="c" * 64, original_filename="lesson.docx"),
            task_input, task_id=task.task_id, plan_id=document.plan_id,
        )
        payload = normalizer.provider.invoke_structured.call_args.kwargs["input_payload"]
        self.assertNotIn("textbook_version", payload["authoritative_metadata"])
        normalizer.provider.invoke_structured.call_args.kwargs["result_validator"](
            normalizer.provider.invoke_structured.return_value.value
        )
        mapped = build_task(task_input, task_id=task.task_id, initial_plan=document)
        self.assertEqual("人教版", mapped.metadata["textbook_version"])

    def test_dataclass_metadata_and_usage_are_preserved(self) -> None:
        task = make_task()
        document = make_document(task, high_quality=True)
        task_input = EngineLessonInput(
            mode=EngineMode.OPTIMIZE,
            subject=task.subject,
            grade=task.grade,
            topic=task.topic,
            duration_minutes=task.duration_minutes,
            optimization_focus=["目标—活动—评价一致性"],
            must_preserve_content=["分数条"],
        )
        output = NormalizationModelOutput(
            metadata_provenance={
                "subject": "user",
                "grade": "user",
                "topic": "user",
                "duration_minutes": "user",
                "textbook_version": "user",
            },
            lesson_plan=document,
            preserved_content_map={"分数条": ["paragraph:2"]},
            raw_locator_coverage={"learning_objectives": ["paragraph:1"]},
        )
        metadata = AgentCallMetadata(
            provider="deepseek",
            model_name="deepseek-v4-flash",
            prompt_id="docx_normalizer_prompt",
            prompt_version="1.0",
            prompt_sha256="b" * 64,
            attempts=2,
        )
        normalizer = DocxNormalizer(make_config())
        normalizer.provider = Mock()
        normalizer.provider.invoke_structured.return_value = ProviderCallResult(
            value=output,
            usage=TokenUsage(input_tokens=321, output_tokens=123),
            estimated_cost=0.0123,
            metadata=metadata,
        )

        normalized = normalizer.normalize(
            RawLessonDocument(
                source_sha256="c" * 64,
                original_filename="lesson.docx",
            ),
            task_input,
            task_id=task.task_id,
            plan_id=document.plan_id,
        )

        self.assertEqual(2, normalized.model_metadata["attempts"])
        self.assertEqual(321, normalized.usage.input_tokens)
        self.assertEqual(0.0123, normalized.estimated_cost)

    def test_projection_includes_preprocessing_usage_without_pipeline_result(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            settings = make_settings(Path(directory))
            registry = RunRegistry(settings)
            record = RunRecord(
                engine_run_id="run-normalizing",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.OPTIMIZE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
                status=EngineRunStatus.RUNNING,
                preprocessing_model_call_count=2,
                preprocessing_input_tokens=321,
                preprocessing_output_tokens=123,
                preprocessing_estimated_cost=0.0123,
            )
            registry.create(record)

            usage = EngineProjection(settings).snapshot(record.engine_run_id).usage

            self.assertEqual(2, usage.model_call_count)
            self.assertEqual(321, usage.input_tokens)
            self.assertEqual(123, usage.output_tokens)
            self.assertEqual(0.0123, usage.estimated_cost)

    def test_failed_normalizer_invocation_keeps_attempt_accounting(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            settings = make_settings(Path(directory))
            registry = RunRegistry(settings)
            record = RunRecord(
                engine_run_id="run-normalizer-failed",
                external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH,
                mode=EngineMode.OPTIMIZE,
                task_id="lesson-one",
                subject="数学",
                grade="高一",
                topic="函数",
            )
            registry.create(record)
            error = ProviderInvocationError(
                "upstream unavailable",
                attempts=3,
                usage=TokenUsage(input_tokens=90, output_tokens=10),
                estimated_cost=0.004,
            )

            _fail(registry, record.engine_run_id, error)

            failed = registry.get(record.engine_run_id)
            self.assertEqual(EngineRunStatus.FAILED, failed.status)
            self.assertEqual(3, failed.preprocessing_model_call_count)
            self.assertEqual(90, failed.preprocessing_input_tokens)
            self.assertEqual(10, failed.preprocessing_output_tokens)
            self.assertEqual(0.004, failed.preprocessing_estimated_cost)
            self.assertIn("upstream unavailable", failed.internal_error_detail)


class GeneratedWordRoundTripTests(unittest.TestCase):
    def test_worker_skips_model_normalizer_for_verified_export(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            settings = make_settings(root)
            registry = RunRegistry(settings)
            record = RunRecord(
                engine_run_id="roundtrip-run", external_job_id=JOB_ID,
                request_sha256=REQUEST_HASH, mode=EngineMode.OPTIMIZE,
                task_id="new-task", subject="数学", grade="五年级", topic="分数的意义",
                original_filename="original.docx",
            )
            registry.create(record)
            source_path = registry.run_dir(record.engine_run_id) / "input" / "original.docx"
            source_path.parent.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph("教学目标：理解分数的意义")
            doc.save(source_path)
            source = make_document(make_task(), high_quality=True)
            normalized = NormalizedLessonInput(
                source_sha256=sha256_file(source_path),
                metadata_provenance={"subject": "user"},
                lesson_plan=source.model_copy(update={"task_id": "new-task"}),
                model_metadata={"provider": "verified_export", "attempts": 0},
            )
            task_input = EngineLessonInput(
                mode=EngineMode.OPTIMIZE, subject="数学", grade="五年级",
                topic="分数的意义", duration_minutes=45,
            )
            with patch("paper4_pipeline.web_api.runner.restore_generated_docx", return_value=normalized), \
                 patch("paper4_pipeline.web_api.runner.DocxNormalizer") as model_normalizer, \
                 patch("paper4_pipeline.web_api.runner._run_pipeline") as run_pipeline:
                execute_optimize(settings.worker_payload(), record.engine_run_id,
                                 task_input.model_dump(mode="json"))
            model_normalizer.assert_not_called()
            run_pipeline.assert_called_once()
            self.assertEqual(0, registry.get(record.engine_run_id).preprocessing_model_call_count)
            self.assertTrue((registry.run_dir(record.engine_run_id) / "normalized" / "initial_plan.json").is_file())

    def test_verified_word_restores_exact_source_without_model(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "artifacts" / "source-run"
            folder.mkdir(parents=True)
            word = folder / "best_lesson_plan.docx"
            doc = Document()
            doc.add_paragraph("分数的意义")
            doc.save(word)
            source = make_document(make_task(), high_quality=True)
            source = source.model_copy(update={
                "metadata": source.metadata.model_copy(update={"subject": "初中数学"})
            })
            exported = folder / "best_lesson_plan.json"
            exported.write_text(json.dumps({
                "schema_version": "paper4-lesson-plan-export-v0.1",
                "run_id": "source-run", "recovery_only": False,
                "lesson_plan": source.model_dump(mode="json"),
            }, ensure_ascii=False), encoding="utf-8")
            (folder / "manifest.json").write_text(json.dumps({
                "run_id": "source-run",
                "artifacts": [
                    {"artifact_id": "best-plan-docx", "status": "ok", "sha256": sha256_file(word)},
                    {"artifact_id": "best-plan-json", "status": "ok", "sha256": sha256_file(exported)},
                ],
            }), encoding="utf-8")
            raw = RawLessonDocument(
                source_sha256=sha256_file(word), original_filename="best_lesson_plan.docx",
            )
            task = make_task()
            requested = EngineLessonInput(
                mode=EngineMode.OPTIMIZE, subject=task.subject,
                grade=task.grade, topic=task.topic,
                duration_minutes=task.duration_minutes,
            )
            restored = restore_generated_docx(
                raw, requested, task_id="new-task", plan_id="new-plan",
                artifacts_root=folder.parent,
            )
            self.assertIsNotNone(restored)
            assert restored is not None
            self.assertEqual("new-task", restored.lesson_plan.task_id)
            self.assertEqual("new-plan", restored.lesson_plan.plan_id)
            self.assertEqual(source.procedure_steps, restored.lesson_plan.procedure_steps)
            self.assertEqual("数学", restored.lesson_plan.metadata.subject)
            self.assertTrue(any("确认等价" in item for item in restored.warnings))
            self.assertEqual(0, restored.model_metadata["attempts"])

            # A changed Word, invalid companion JSON, or changed task identity
            # must take the normal importer path. Extra user constraints do not
            # invalidate the verified document/JSON pair.
            self.assertIsNone(restore_generated_docx(
                raw.model_copy(update={"source_sha256": "0" * 64}), requested,
                task_id="new-task", plan_id="new-plan", artifacts_root=folder.parent,
            ))
            self.assertIsNone(restore_generated_docx(
                raw, requested.model_copy(update={"topic": "另一课题"}),
                task_id="new-task", plan_id="new-plan", artifacts_root=folder.parent,
            ))
            preserved = restore_generated_docx(
                raw, requested.model_copy(update={"must_preserve_content": ["分数条"]}),
                task_id="new-task", plan_id="new-plan", artifacts_root=folder.parent,
            )
            self.assertIsNotNone(preserved)
            assert preserved is not None
            self.assertEqual({"分数条": []}, preserved.preserved_content_map)
            self.assertTrue(any("人工核对" in warning for warning in preserved.warnings))
            changed_textbook = restore_generated_docx(
                raw, requested.model_copy(update={"textbook_version": "新版教材"}),
                task_id="new-task", plan_id="new-plan", artifacts_root=folder.parent,
            )
            self.assertIsNotNone(changed_textbook)
            assert changed_textbook is not None
            self.assertEqual("新版教材", changed_textbook.lesson_plan.metadata.textbook_version)
            self.assertEqual("user", changed_textbook.metadata_provenance["textbook_version"])
            exported.write_text("{}", encoding="utf-8")
            self.assertIsNone(restore_generated_docx(
                raw, requested, task_id="new-task", plan_id="new-plan",
                artifacts_root=folder.parent,
            ))


class EngineApiContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.executor = RecordingExecutor()
        settings = make_settings(self.root)
        manager = EngineManager(settings, executor=self.executor)
        self.client_context = TestClient(create_app(settings, manager))
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        self.temporary.cleanup()

    def test_frozen_internal_contract_fixtures_are_valid(self) -> None:
        fixtures = PACKAGE_ROOT / "tests" / "fixtures" / "web_contracts"
        request = CreateRunRequest.model_validate_json(
            (fixtures / "engine-generate-request.json").read_text(encoding="utf-8")
        )
        accepted = RunAccepted.model_validate_json(
            (fixtures / "engine-run-accepted.json").read_text(encoding="utf-8")
        )
        self.assertEqual("01J00000000000000000000000", request.external_job_id)
        self.assertEqual("官能团与有机物性质", request.task.topic)
        self.assertEqual("01J00000000000000000000000", accepted.external_job_id)

    def test_generate_requires_internal_token_and_returns_202(self) -> None:
        denied = self.client.post(
            "/internal/v1/runs/generate", json=request_payload()
        )
        self.assertEqual(401, denied.status_code)
        self.assertEqual("ENGINE_UNAUTHORIZED", denied.json()["code"])
        self.assertEqual("application/problem+json", denied.headers["content-type"])

        response = self.client.post(
            "/internal/v1/runs/generate",
            json=request_payload(),
            headers={"X-Engine-Token": "test-token"},
        )
        self.assertEqual(202, response.status_code)
        self.assertEqual(JOB_ID, response.json()["external_job_id"])
        self.assertEqual(1, len(self.executor.calls))

        repeated = self.client.post(
            "/internal/v1/runs/generate",
            json=request_payload(),
            headers={"X-Engine-Token": "test-token"},
        )
        self.assertEqual(response.json()["engine_run_id"], repeated.json()["engine_run_id"])
        self.assertEqual(1, len(self.executor.calls))

    def test_optimize_streams_original_file_into_run_directory(self) -> None:
        document = Document()
        document.add_paragraph("教学目标：理解一次函数")
        path = self.root / "upload.docx"
        document.save(path)
        payload = request_payload(mode="optimize")
        response = self.client.post(
            "/internal/v1/runs/optimize",
            data={"request": json.dumps(payload, ensure_ascii=False)},
            files={
                "document": (
                    "原教案.docx",
                    path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers={"X-Engine-Token": "test-token"},
        )
        self.assertEqual(202, response.status_code, response.text)
        run_id = response.json()["engine_run_id"]
        saved = self.root / "state" / "runs" / run_id / "input" / "original.docx"
        self.assertTrue(saved.is_file())
        self.assertEqual(1, len(self.executor.calls))


if __name__ == "__main__":
    unittest.main()
