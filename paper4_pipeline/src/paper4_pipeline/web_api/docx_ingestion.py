"""Security-first DOCX inspection and format-agnostic content extraction."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from paper4_pipeline.web_api.schemas import (
    RawBlock,
    RawLessonDocument,
    RawParagraph,
    RawTable,
    RawTableCell,
)
from paper4_pipeline.web_api.settings import EngineSettings


OOXML_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_REQUIRED_PARTS = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
_FORBIDDEN_SUFFIXES = ("vbaProject.bin", "vbaData.xml")
_FORBIDDEN_PARTS = {"EncryptionInfo", "EncryptedPackage"}
_RELATIONSHIP_LABELS = {
    "hyperlink": "外部超链接",
    "image": "外链图片",
    "attachedtemplate": "外部模板",
}


class DocxIngestionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_member_name(name: str) -> bool:
    if not name or "\\" in name or name.startswith(("/", "\\")):
        return False
    pure = PurePosixPath(name)
    return not pure.is_absolute() and ".." not in pure.parts


def _external_relation_kind(relation_type: str) -> str:
    """Classify a relationship without ever resolving or fetching its target."""

    suffix = relation_type.rstrip("/").rsplit("/", 1)[-1].casefold()
    return _RELATIONSHIP_LABELS.get(suffix, "外部关系")


def inspect_docx(path: Path, settings: EngineSettings) -> list[str]:
    if path.suffix.lower() != ".docx":
        raise DocxIngestionError("INVALID_DOCX", "仅支持 .docx 文件")
    size = path.stat().st_size
    if size <= 0:
        raise DocxIngestionError("INVALID_DOCX", "Word 文件为空")
    if size > settings.max_upload_bytes:
        raise DocxIngestionError("DOCX_TOO_LARGE", "Word 文件超过 20 MiB 限制")
    with path.open("rb") as handle:
        if handle.read(4) != b"PK\x03\x04":
            raise DocxIngestionError("INVALID_DOCX", "文件不是有效的 OOXML ZIP 包")

    warnings: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > settings.max_zip_entries:
                raise DocxIngestionError("UNSAFE_DOCX", "Word 包含过多内部文件")
            names = {item.filename for item in infos}
            lower_names = {item.lower() for item in names}
            missing = _REQUIRED_PARTS - names
            if missing:
                raise DocxIngestionError(
                    "INVALID_DOCX", "Word 缺少必要 OOXML 组成部分"
                )
            expanded = 0
            for info in infos:
                if not _safe_member_name(info.filename):
                    raise DocxIngestionError("UNSAFE_DOCX", "Word 包含不安全路径")
                lowered = info.filename.lower()
                if lowered in {item.lower() for item in _FORBIDDEN_PARTS} or lowered.endswith(
                    tuple(item.lower() for item in _FORBIDDEN_SUFFIXES)
                ):
                    raise DocxIngestionError("UNSAFE_DOCX", "不支持加密或含宏的 Word")
                if lowered.startswith("word/embeddings/"):
                    raise DocxIngestionError("UNSAFE_DOCX", "不支持含 OLE 嵌入对象的 Word")
                if info.file_size > settings.max_zip_entry_bytes:
                    raise DocxIngestionError("UNSAFE_DOCX", "Word 内部单项过大")
                expanded += info.file_size
                if expanded > settings.max_zip_expanded_bytes:
                    raise DocxIngestionError("UNSAFE_DOCX", "Word 展开后体积过大")
                compressed = max(info.compress_size, 1)
                ratio = info.file_size / compressed
                if info.file_size > 1024 * 1024 and ratio > settings.max_compression_ratio:
                    raise DocxIngestionError("UNSAFE_DOCX", "Word 压缩比异常")
            external_relation_counts: dict[str, int] = {}
            for name in names:
                if not name.endswith(".rels"):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as exc:
                    raise DocxIngestionError("INVALID_DOCX", "Word 关系文件损坏") from exc
                for relation in root.iter():
                    if not relation.tag.endswith("Relationship"):
                        continue
                    if relation.attrib.get("TargetMode", "").lower() == "external":
                        # OOXML uses external relationships for ordinary links,
                        # linked header images and attached authoring templates.
                        # Text extraction never dereferences their targets, so
                        # rejecting the whole lesson loses valid classroom text
                        # without adding security. Embedded OLE, macros and
                        # encrypted packages remain hard failures above.
                        label = _external_relation_kind(
                            relation.attrib.get("Type", "")
                        )
                        external_relation_counts[label] = (
                            external_relation_counts.get(label, 0) + 1
                        )
            for label, count in sorted(external_relation_counts.items()):
                warnings.append(
                    f"原 Word 含 {count} 个{label}；系统不会访问其目标，"
                    "仅提取文件中已有的正文和表格"
                )
            if any(name.startswith("word/comments") for name in lower_names):
                warnings.append("原 Word 含批注，MVP 不保证提取批注内容")
            if any(name.startswith("word/media/") for name in lower_names):
                warnings.append("原 Word 含图片，MVP 不执行 OCR，请人工核对图片信息")
            if "word/footnotes.xml" in lower_names or "word/endnotes.xml" in lower_names:
                warnings.append("原 Word 含脚注或尾注，MVP 仅提取正文与表格")
    except zipfile.BadZipFile as exc:
        raise DocxIngestionError("INVALID_DOCX", "Word ZIP 结构损坏") from exc
    return warnings


def _clean_text(value: str) -> str:
    return re.sub(r"[\t\r ]+", " ", value.replace("\u00a0", " ")).strip()


def extract_docx(
    path: Path,
    *,
    original_filename: str,
    settings: EngineSettings,
) -> RawLessonDocument:
    warnings = inspect_docx(path, settings)
    try:
        from docx import Document
    except ImportError as exc:
        raise DocxIngestionError(
            "DOCX_DEPENDENCY_MISSING", "服务未安装 python-docx"
        ) from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocxIngestionError("INVALID_DOCX", "无法读取 Word 正文") from exc

    from docx.table import Table

    paragraphs: list[RawParagraph] = []
    tables: list[RawTable] = []
    ordered_blocks: list[RawBlock] = []
    paragraph_index = 0
    table_index = 0
    for element in document.iter_inner_content():
        if not isinstance(element, Table):
            paragraph_index += 1
            text = _clean_text(element.text)
            if not text:
                continue
            try:
                style = element.style.name or ""
            except Exception:
                style = ""
            locator = f"paragraph:{paragraph_index}"
            paragraphs.append(RawParagraph(locator=locator, style=style, text=text))
            ordered_blocks.append(RawBlock(
                locator=locator, kind="paragraph", text=text, style=style,
            ))
            continue
        table_index += 1
        table = element
        cells: list[RawTableCell] = []
        max_columns = 0
        for row_index, row in enumerate(table.rows):
            max_columns = max(max_columns, len(row.cells))
            row_parts: list[str] = []
            # ``python-docx`` repeats the same XML cell once per occupied grid
            # column for horizontally merged cells. Treat it as one source
            # cell; otherwise a common multi-column lesson-plan template can
            # multiply the same paragraph many times and distort both token
            # cost and model interpretation.
            seen_xml_cells: set[int] = set()
            for column_index, cell in enumerate(row.cells):
                xml_identity = id(cell._tc)
                if xml_identity in seen_xml_cells:
                    continue
                seen_xml_cells.add(xml_identity)
                text = _clean_text("\n".join(p.text for p in cell.paragraphs))
                if not text:
                    continue
                row_parts.append(f"列{column_index + 1}: {text}")
                cells.append(
                    RawTableCell(
                        locator=(
                            f"table:{table_index}/row:{row_index + 1}/"
                            f"column:{column_index + 1}"
                        ),
                        row=row_index,
                        column=column_index,
                        text=text,
                    )
                )
            if row_parts:
                ordered_blocks.append(RawBlock(
                    locator=f"table:{table_index}/row:{row_index + 1}",
                    kind="table_row", text=" | ".join(row_parts),
                ))
        tables.append(
            RawTable(
                locator=f"table:{table_index}",
                rows=len(table.rows),
                columns=max_columns,
                cells=cells,
            )
        )

    if not paragraphs and not any(table.cells for table in tables):
        raise DocxIngestionError("DOCX_EMPTY_CONTENT", "Word 中没有可提取的正文或表格")
    raw = RawLessonDocument(
        source_sha256=sha256_file(path),
        original_filename=original_filename,
        paragraphs=paragraphs,
        tables=tables,
        ordered_blocks=ordered_blocks,
        warnings=warnings,
    )
    if raw.character_count() > settings.max_normalizer_characters:
        raise DocxIngestionError(
            "DOCX_CONTENT_TOO_LARGE", "Word 可读文字过多，超过结构化处理上限"
        )
    return raw
