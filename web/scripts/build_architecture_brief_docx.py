"""Build a standalone Lessongen architecture brief and its diagrams."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


WEB_ROOT = Path(__file__).resolve().parents[1]
PY_ROOT = WEB_ROOT.parent / "paper4_pipeline"
OUTPUT = WEB_ROOT / "docs" / "Lessongen_核心架构与Web实现.docx"
IMAGES = WEB_ROOT / "docs" / "ui-audit"
DIAGRAM = WEB_ROOT / "docs" / "ui-audit" / "architecture-current.png"
AGENT_DIAGRAM = WEB_ROOT / "docs" / "ui-audit" / "agent-routing-current.png"


def draw_architecture() -> None:
    """Render a concise schematic of the implemented request and execution path."""
    canvas = Image.new("RGB", (1800, 1250), "white")
    draw = ImageDraw.Draw(canvas)
    regular = r"C:\Windows\Fonts\msyh.ttc"
    bold = r"C:\Windows\Fonts\msyhbd.ttc"
    title_font = ImageFont.truetype(bold, 46)
    body_font = ImageFont.truetype(regular, 33)
    note_font = ImageFont.truetype(regular, 27)
    ink = "#18323B"
    line = "#6B9197"

    def centered(text: str, center_x: int, top: int, font, fill=ink) -> None:
        bounds = draw.textbbox((0, 0), text, font=font)
        draw.text((center_x - (bounds[2] - bounds[0]) / 2, top), text, font=font, fill=fill)

    def box(coords, heading: str, lines: list[str], fill: str) -> None:
        draw.rounded_rectangle(coords, radius=22, fill=fill, outline="#A9C6C9", width=3)
        cx = (coords[0] + coords[2]) // 2
        centered(heading, cx, coords[1] + 24, title_font)
        top = coords[1] + 85
        for item in lines:
            centered(item, cx, top, body_font)
            top += 47

    def arrow(points, *, color=line) -> None:
        draw.line(points, fill=color, width=8, joint="curve")
        x0, y0 = points[-2]
        x1, y1 = points[-1]
        if x1 > x0:
            head = [(x1, y1), (x1 - 23, y1 - 14), (x1 - 23, y1 + 14)]
        elif x1 < x0:
            head = [(x1, y1), (x1 + 23, y1 - 14), (x1 + 23, y1 + 14)]
        else:
            head = [(x1, y1), (x1 - 14, y1 - 23), (x1 + 14, y1 - 23)]
        draw.polygon(head, fill=color)

    box((70, 45, 560, 250), "Vue 前端", ["输入与上传", "进度 预览 下载"], "#EDF6F5")
    box((655, 45, 1145, 250), "Spring Boot", ["公开 API 任务调度", "MySQL 与本地文件"], "#F0F4F7")
    box((1240, 45, 1730, 250), "FastAPI", ["内部鉴权 作业登记", "单进程模型工作者"], "#EDF6F5")
    arrow([(565, 147), (650, 147)])
    arrow([(1150, 147), (1235, 147)])
    centered("REST", 608, 267, note_font)
    centered("X-Engine-Token", 1195, 267, note_font)

    arrow([(1485, 250), (1485, 326), (505, 326), (505, 370)])
    arrow([(1485, 326), (1295, 326), (1295, 370)])
    box((180, 370, 830, 555), "生成入口", ["任务信息 → Design Architect", "Writer 生成初稿 v0"], "#F4F8F8")
    box((970, 370, 1620, 555), "优化入口", ["DOCX 安全抽取", "模型规范化为原稿 v0"], "#F4F8F8")
    arrow([(505, 555), (505, 685)])
    arrow([(1295, 555), (1295, 685)])

    box(
        (180, 685, 1620, 955),
        "LangGraph 共用质量闭环",
        [
            "Judge → Subject Pedagogy Alignment 三类 Critic",
            "Validator → Rewriter → Verifier → 复评或停靠",
            "统一 LessonPlanDocument 版本与程序约束",
        ],
        "#EAF3F2",
    )
    arrow([(505, 955), (505, 1065)])
    arrow([(1295, 955), (1295, 1065)])
    box((180, 1065, 830, 1225), "模型服务", ["DeepSeek V4 Flash"], "#F0F4F7")
    box((970, 1065, 1620, 1225), "运行产物", ["JSON Markdown Word trace"], "#F0F4F7")
    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(DIAGRAM, optimize=True)


def draw_agent_architecture() -> None:
    """Diagram the exact implemented LangGraph nodes and conditional edges."""
    canvas = Image.new("RGB", (1800, 2070), "white")
    draw = ImageDraw.Draw(canvas)
    regular = r"C:\Windows\Fonts\msyh.ttc"
    bold = r"C:\Windows\Fonts\msyhbd.ttc"
    title = ImageFont.truetype(bold, 42)
    body = ImageFont.truetype(regular, 31)
    small = ImageFont.truetype(regular, 27)
    ink = "#18323B"
    teal = "#658D94"
    stop = "#A67553"

    def centered(text: str, cx: int, y: int, font, color=ink) -> None:
        rect = draw.textbbox((0, 0), text, font=font)
        draw.text((cx - (rect[2] - rect[0]) / 2, y), text, font=font, fill=color)

    def box(coords, heading: str, lines: list[str], fill="#EFF7F6", edge="#A9C6C9") -> None:
        draw.rounded_rectangle(coords, radius=22, fill=fill, outline=edge, width=3)
        cx = (coords[0] + coords[2]) // 2
        centered(heading, cx, coords[1] + 17, title)
        y = coords[1] + 70
        for text in lines:
            centered(text, cx, y, body)
            y += 44

    def arrow(points, color=teal) -> None:
        draw.line(points, fill=color, width=7, joint="curve")
        x0, y0 = points[-2]
        x1, y1 = points[-1]
        if x1 > x0:
            head = [(x1, y1), (x1 - 22, y1 - 14), (x1 - 22, y1 + 14)]
        elif x1 < x0:
            head = [(x1, y1), (x1 + 22, y1 - 14), (x1 + 22, y1 + 14)]
        else:
            head = [(x1, y1), (x1 - 14, y1 - 22), (x1 + 14, y1 - 22)]
        draw.polygon(head, fill=color)

    box((100, 45, 800, 215), "生成入口", ["Design Architect → Writer", "形成结构化初稿 v0"])
    box((1000, 45, 1700, 215), "优化入口", ["DOCX 抽取 → 模型规范化", "以用户原稿形成 v0"])
    arrow([(450, 215), (450, 267), (900, 267), (900, 315)])
    arrow([(1350, 215), (1350, 267), (900, 267), (900, 315)])

    box((440, 315, 1360, 475), "Judge Agent", ["八维内部评价 风险与未解决项", "程序另外选出当前最佳候选"])
    arrow([(900, 475), (900, 550)])
    box((440, 550, 1360, 705), "评估路由  程序节点", ["质量门槛 轮数 预算 停滞 振荡 回归", "继续批评或停靠"] , fill="#EDF2F6")
    box((1450, 550, 1755, 705), "Finalize", ["完成 复核", "或失败"], fill="#FAF4EE", edge="#D9BDAB")
    arrow([(1360, 625), (1450, 625)], stop)

    draw.line([(900, 705), (900, 775)], fill=teal, width=7)
    draw.line([(310, 775), (1490, 775)], fill=teal, width=7)
    for x in (310, 900, 1490):
        arrow([(x, 775), (x, 825)])
    box((60, 825, 560, 1010), "Subject Critic", ["学科事实 教材边界", "误概念与知识顺序"])
    box((650, 825, 1150, 1010), "Pedagogy Critic", ["教学逻辑 学习支架", "参与和课堂可行性"])
    box((1240, 825, 1740, 1010), "Alignment Critic", ["课程目标 活动产出", "评价证据链对齐"])
    draw.line([(310, 1010), (310, 1060)], fill=teal, width=7)
    draw.line([(900, 1010), (900, 1060)], fill=teal, width=7)
    draw.line([(1490, 1010), (1490, 1060)], fill=teal, width=7)
    draw.line([(310, 1060), (1490, 1060)], fill=teal, width=7)
    arrow([(900, 1060), (900, 1115)])

    box((350, 1115, 1450, 1295), "意见聚合与 Validator Agent", ["汇总同轮意见 核查依据和可执行性", "接受 拒绝 合并 延期  每轮最多接受五条"])
    arrow([(900, 1295), (900, 1385)])
    box((440, 1385, 1360, 1520), "验证路由  程序节点", ["有已接受意见且预算允许才进入改写"], fill="#EDF2F6")
    box((1450, 1385, 1755, 1550), "Finalize", ["无可执行意见", "或预算不足"], fill="#FAF4EE", edge="#D9BDAB")
    arrow([(1360, 1452), (1450, 1452)], stop)
    arrow([(900, 1520), (900, 1610)])
    box((400, 1610, 1400, 1785), "Rewriter Agent", ["仅按接受意见做必要修改并记录映射", "失败则保留既有版本并转人工复核"])
    box((1490, 1610, 1770, 1785), "Finalize", ["改写失败", "人工复核"], fill="#FAF4EE", edge="#D9BDAB")
    arrow([(1400, 1697), (1490, 1697)], stop)
    arrow([(900, 1785), (900, 1850)])
    box((350, 1850, 1450, 2030), "Verifier  程序节点", ["核查改动是否落实 检测质量回归", "记录结果后回到 Judge 评估新版本"] , fill="#EDF2F6")
    arrow([(350, 1930), (25, 1930), (25, 395), (440, 395)])
    centered("下一轮复评", 181, 1800, small, teal)
    AGENT_DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(AGENT_DIAGRAM, optimize=True)


def set_font(style, name: str, size: int, *, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_table(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers: tuple[str, ...], rows: list[tuple[str, ...]], widths: tuple[float, ...]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = title
        shade(cell, "E9F1F0")
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9.3)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.1)
    border_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_p(doc, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_shot(doc, title: str, filename: str, *, width: float = 5.65, start_new_page: bool = True):
    heading = doc.add_heading(title, level=1)
    if start_new_page:
        heading.paragraph_format.page_break_before = True
    image_path = IMAGES / filename
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run("界面验收截图；页面内容使用固定契约示例数据，不代表该截图来自付费模型运行。")


def main() -> None:
    draw_architecture()
    draw_agent_architecture()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.70)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.34)

    styles = doc.styles
    set_font(styles["Normal"], "Microsoft YaHei", 9.7)
    styles["Normal"].paragraph_format.line_spacing = 1.22
    styles["Normal"].paragraph_format.space_after = Pt(6)
    set_font(styles["Title"], "Microsoft YaHei", 18, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(8)
    set_font(styles["Heading 1"], "Microsoft YaHei", 12.5, bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(13)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    set_font(styles["Heading 2"], "Microsoft YaHei", 10.5, bold=True)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    set_font(styles["Caption"], "Microsoft YaHei", 8)
    styles["Caption"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    title.add_run("Paper4 教案生成与优化系统架构及 Web 实现")
    add_p(doc, "截至 2026 年 9 月 16 日", "Subtitle")
    add_p(
        doc,
        "当前系统已形成真实模型驱动的教案生成与优化闭环，并接入可使用的 Web 工作台。本文只保留关键流程、工程边界和已验证结果，供组会汇报与后续协作使用。系统仍属于研究原型：内部评分不能替代教师评价，Paper3 课堂模拟、对抗机制和人的实验尚未完成。",
    )

    doc.add_heading("当前完成范围", level=1)
    add_table(
        doc,
        ("方向", "已经实现", "当前边界"),
        [
            ("Paper4 核心", "真实 Agent 生成、三类批评、裁决、改写、复评与版本导出", "尚未证明课堂教学效果"),
            ("Web 作品", "生成教案、上传 Word 优化、任务进度、结果预览与下载", "目前只接入 Paper4 两条路径，不是四个 Paper 的完整集成"),
            ("后续研究", "保留角色、知识、Prompt、版本和 trace 的可比较接口", "Paper3 联动、角色辩论、权威 RAG 与真人实验待开展"),
        ],
        (1.08, 3.05, 2.73),
    )
    add_p(
        doc,
        "编号说明：Paper#4 是研究分工中的多智能体教案撰写方向，对应导师产品构想中的“功能二”；产品“功能四”属于 LLM 与人机交互方向。",
    )

    doc.add_heading("核心运行架构", level=1)
    add_p(doc, "两种入口进入同一套结构化教案与质量迭代流程：")
    add_p(doc, "生成入口：教学信息 → Design Architect 选择设计路线 → Writer 生成初稿 v0。", "List Bullet")
    add_p(doc, "优化入口：上传 DOCX → 安全检查与内容抽取 → 结构化规范化 → 原稿 v0。", "List Bullet")
    add_p(
        doc,
        "共同闭环：硬规则检查 → Judge 内部评价及程序路由 → Subject、Pedagogy、Alignment 三类 Critic 独立提出意见 → Validator 去重、合并与裁决 → Rewriter 定向修改 → Verifier 检查修改是否落实及是否回归 → 再评价或停止。",
    )
    add_p(
        doc,
        "这里的 Agent 负责生成、诊断和语义判断；程序负责 JSON Schema、身份与引用约束、课时校验、预算、停止条件和最终版本选择。模型的“建议通过”不能直接覆盖程序判断。",
    )
    add_table(
        doc,
        ("角色或组件", "核心职责"),
        [
            ("Design Architect 与 Writer", "先选择教学设计路线，再按统一教案结构形成初稿；优化模式跳过这两步。"),
            ("三类 Critic", "分别审查学科准确性、教学法与学习支持、目标活动评价对齐；本轮互不读取对方输出。"),
            ("Validator", "检查意见是否有依据、相关且可执行，处理重复、冲突、接受、拒绝、合并和延期。"),
            ("Rewriter 与 Verifier", "按已接受意见做必要修改，并核查修改映射、未解决项及质量回退。"),
            ("Judge 与程序路由", "记录八维内部评分；程序结合硬规则、风险、轮数和调用预算决定继续或停靠。"),
        ],
        (1.72, 5.14),
    )

    doc.add_heading("数据和可追溯性", level=1)
    add_p(
        doc,
        "教案使用 LessonPlanDocument 作为统一 JSON 合同，并按版本保存。每轮保留批评意见、Validator 决策、修改映射、评分、路由、Prompt 版本、模型用量和 trace；最佳版本可导出 JSON、Markdown、Word。失败或需人工复核时保留状态与可恢复结果，不冒充质量通过。",
    )
    add_p(
        doc,
        "当前知识主要来自任务输入和可配置的静态知识包，尚非具备检索、证据片段引用与版本治理的完整 RAG。Judge 八维分数用于系统内部比较，不是经教师验证的教学成效指标。",
    )

    doc.add_heading("Web 实现", level=1)
    add_p(
        doc,
        "Vue 页面负责输入、进度和结果展示；Spring Boot 提供公开 API、MySQL 任务记录、异步编排和文件下载；FastAPI 将请求交给现有 PipelineService 与 LangGraph。浏览器不接触模型密钥或 Python 内部接口。",
    )
    add_table(
        doc,
        ("页面功能", "已实现行为"),
        [
            ("生成教案", "填写科目、年级、课题等信息后创建任务；可补充教材、课标、学情、资源与设计要求。"),
            ("优化教案", "上传 .docx，安全抽取并结构化原稿；支持填写优化重点与必须保留的内容，使用同一迭代图。"),
            ("过程与结果", "显示真实阶段和轮次，支持断线轮询；查看教案、内部评分、修改与复核信息，下载产物。"),
            ("历史任务", "按模式和状态查看既有任务；任务和文件元数据保存在 MySQL，文件保存在受控存储。"),
        ],
        (1.36, 5.50),
    )
    add_p(
        doc,
        "运行数据目录由环境变量配置；Web 的上传、任务状态和下载依赖本机服务运行，当前版本尚未进行多用户生产部署。",
    )

    doc.add_heading("系统总体架构", level=1)
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_after = Pt(3)
    figure.add_run().add_picture(str(DIAGRAM), width=Inches(6.52))
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run("图一 当前已实现的 Web 到模型运行链路")
    add_p(
        doc,
        "浏览器只请求 Spring Boot。Java 后端先保存任务及上传文件，再由定时调度向 FastAPI 提交作业；随后轮询引擎状态和事件，将进度、结果及可下载产物提供给前端。MySQL 保存任务、事件和文件元数据，不保存完整教案生成过程。",
    )
    add_p(
        doc,
        "Python 引擎使用单工作进程执行真实模型任务。生成模式先由 Design Architect 与 Writer 形成初稿；优化模式先安全抽取 Word，再由模型规范化原稿。两种模式从初稿起共用 Judge、三类 Critic、Validator、Rewriter、Verifier 的迭代图；模型调用通过 OpenAI 兼容接口访问 DeepSeek V4 Flash。",
    )
    add_p(
        doc,
        "本机文件目录保存上传原件、引擎状态、各轮 trace 以及 JSON、Markdown、Word 导出物。模型密钥只配置在 Python 端；Java 与 Python 之间使用 X-Engine-Token 鉴权，Vue 不持有这两类密钥。图中为当前实现，不包含尚未接入的 Paper3 课堂模拟或多人部署。",
    )

    doc.add_heading("Agent 运行架构与路由", level=1)
    agent_figure = doc.add_paragraph()
    agent_figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agent_figure.paragraph_format.space_after = Pt(3)
    agent_figure.add_run().add_picture(str(AGENT_DIAGRAM), width=Inches(6.48))
    agent_caption = doc.add_paragraph(style="Caption")
    agent_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agent_caption.add_run("图二 当前 LangGraph 节点与条件路由 右侧为停靠 左侧回路为下一轮")

    agent_details_heading = doc.add_heading("Agent 节点与路由说明", level=1)
    agent_details_heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ("节点", "本步处理与产出", "下一跳"),
        [
            ("Design Architect", "仅在生成模式读取任务和设计知识，提出并选择教学设计路线，输出 DesignBlueprint。", "进入 Writer。优化模式跳过。"),
            ("Writer 与初稿", "Writer 依据蓝图和任务生成 LessonPlanDocument v0，并通过结构及硬规则检查；优化模式把上传的 DOCX 抽取、规范化为 v0。", "v0 进入 Judge。"),
            ("Judge", "对当前版本给出八维内部评分、风险和未解决项；程序同时选出截至当前的最佳候选版本。模型的推荐动作仅作为证据。", "进入评估路由，不由 Judge 自行停机。"),
            ("评估路由", "程序核查硬规则、风险、总分和最低维度；再检查轮次、调用和 token 预算、成本、耗时、停滞、振荡与回归。", "未触发停止才进入三类 Critic；否则停靠。"),
            ("三类 Critic", "Subject 审学科事实与教材边界；Pedagogy 审教学过程和学生支持；Alignment 审目标、活动、产出、评价对齐。三者使用不同知识包，对同一版本分别给出 CritiqueItem。", "程序聚合同轮意见后交 Validator。"),
            ("Validator", "逐条核查依据、相关性和可执行性，决定接受、拒绝、合并或延期；程序再执行证据限制与每轮最多五条接受意见的上限。", "输出有状态的意见批次，进入验证路由。"),
            ("验证路由", "程序确认是否存在已接受意见，以及预算是否足够覆盖改写和再次评价。", "满足条件才进入 Rewriter；否则停靠。"),
            ("Rewriter", "只使用已接受意见，对每条意见记录修改或未解决映射；新版本必须保持任务身份并通过硬规则。", "成功生成 v1、v2 等版本并进入 Verifier；失败保留旧版本转人工复核。"),
            ("Verifier", "程序核查声称的修改是否落在目标字段，并检查曾解决的问题是否回归；它不是一次新的 LLM Agent 调用。", "回到 Judge 评价新版本。"),
            ("Finalize", "按规则从安全候选中选最佳版本，不默认选择最新版本；明确区分完成、需人工复核和失败。", "输出结果、trace 及可用的 JSON、Markdown、Word。"),
        ],
        (1.36, 3.78, 1.72),
    )
    role_limits_heading = doc.add_heading("角色配置与停止条件", level=2)
    role_limits_heading.paragraph_format.page_break_before = True
    add_p(
        doc,
        "当前配置的质量停靠条件是硬规则通过、无高风险、总分至少 8.0 且八维最低分至少 7.0；最多改写三轮。调用、总 token、估计费用和运行时间上限分别为 42 次、36 万、5 美元和 1800 秒。Judge 的八维为课程对齐、知识准确性、教学逻辑、课堂可行性、差异化教学、学生参与、评价设计、语言与格式。评分用于管线内部比较，尚未通过教师实验验证。",
    )
    add_p(
        doc,
        "这些角色目前都调用 deepseek-v4-flash，差异由角色 Prompt、可见输入、知识源和模型参数形成。三类 Critic 在代码中依次调用，但每人只看同一稿件及此前轮次的意见，不读取本轮其他 Critic 的输出。路由、意见聚合、Verifier、硬规则和最终版本选择由程序执行；当前没有角色辩论、Paper3 模拟回流或每轮必到场的人工 Reviewer。",
    )

    doc.add_heading("验证与下一步", level=1)
    validation_summary = add_p(
        doc,
        "已完成一次真实生成和一次真实 Word 优化闭环，两次均使用 deepseek-v4-flash 并产出 Word。Python 回归为 110 项通过；Web 前端 7 条浏览器端到端测试通过。Java 最近一次复测为 18 项通过、1 项 MySQL Testcontainers 用例因 Docker 未运行而跳过；该容器用例曾在 MySQL 8.0.44 环境通过。",
    )
    validation_summary.paragraph_format.keep_with_next = True
    add_p(
        doc,
        "下一阶段先固定任务集与人工评价量表，比较单轮和多轮、Critic 与知识配置差异；再与 Paper3 对齐模拟反馈格式，验证其能否作为后验意见进入教案修改。师范生与中小学教师的真实使用反馈需要单独设计实验，不能用当前两次工程运行代替。",
    )

    gallery_heading = doc.add_heading("Web 界面截图", level=1)
    gallery_heading.paragraph_format.page_break_before = True
    add_p(doc, "以下为界面验收截图，展示页面能力；其中的页面数据使用固定契约示例。")
    add_shot(doc, "工作台首页", "01-dashboard-desktop.png", width=4.75, start_new_page=False)
    add_shot(doc, "生成教案输入", "02-generate-desktop.png", width=5.45)
    add_shot(doc, "上传 Word 优化", "03-optimize-desktop.png", width=4.95)
    add_shot(doc, "教案结果与下载", "04-result-desktop.png", width=5.30)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
